from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
DESKTOP = Path.home() / "Desktop"
OUTPUT = DESKTOP / "cali.docx"


def set_run_font(run, size=11, bold=False, ascii_font="Times New Roman", east_font="SimSun"):
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_font)
    run.font.size = Pt(size)
    run.bold = bold


def add_para(doc, text, size=11, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_code_block(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(code.rstrip())
    set_run_font(run, size=9.5, ascii_font="Consolas", east_font="Consolas")
    return p


def get_snippet(path_str, start, end):
    path = ROOT / path_str
    lines = path.read_text(encoding="utf-8").splitlines()
    return "\n".join(lines[start - 1 : end])


FORMULAS = [
    {
        "title": "公式1 马尔可夫决策过程定义",
        "formula": "M = (S, A, P, R, γ)",
        "role": "这条公式定义了整篇论文处理的强化学习任务。它不是训练损失，而是问题本身的数学描述：智能体在状态空间中观察环境，采取动作，环境按概率转移，并返回奖励。",
        "why": "论文后面的策略优化、蒸馏、重置和持续学习，都是在这个统一的 MDP 框架上进行的。如果没有这个定义，后续所有公式都没有共同的任务背景。",
        "symbols": [
            ("M", "整个马尔可夫决策过程，也就是一个强化学习任务。"),
            ("S", "状态空间，表示环境所有可能状态的集合。"),
            ("A", "动作空间，表示智能体可以采取的动作集合。"),
            ("P", "状态转移概率，表示在某状态执行某动作后转移到下一个状态的规律。"),
            ("R", "奖励函数，表示环境对当前行为的即时反馈。"),
            ("γ", "折扣因子，范围通常在 0 到 1 之间；越接近 1，说明越重视未来长期回报。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "599-607",
                "intro": "代码中没有把 M=(S,A,P,R,γ) 原样写成一行，但它体现在环境创建和策略网络初始化中：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 599, 607),
            }
        ],
    },
    {
        "title": "公式2 TD误差",
        "formula": "δ_t = r_t + γ V(s_{t+1}) (1 - d_t) - V(s_t)",
        "role": "TD 误差衡量“当前价值估计错了多少”。如果 δ_t 很大，说明当前策略对这个状态的判断不准，往往意味着这里更难学、也更值得关注。",
        "why": "你的论文里，TD 误差不仅用于基础 PPO 更新，还被继续用于动态 α 调节、困难状态优先蒸馏和自适应重置，因此它是全篇最核心的反馈信号之一。",
        "symbols": [
            ("δ_t", "t 时刻的 TD 误差。"),
            ("r_t", "t 时刻执行动作后获得的即时奖励。"),
            ("γ", "折扣因子，用来衡量下一时刻价值对当前的重要程度。"),
            ("V(s_{t+1})", "下一状态的价值估计。"),
            ("d_t", "终止标记；若 episode 结束，d_t=1，否则 d_t=0。"),
            ("1 - d_t", "如果 episode 已结束，就把下一状态价值项关掉。"),
            ("V(s_t)", "当前状态的价值估计。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "413-420",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 413, 420),
            }
        ],
    },
    {
        "title": "公式3 PPO策略损失",
        "formula": "L_actor = - E[min(r_t A_t, clip(r_t, 1-ε, 1+ε) A_t)]",
        "role": "这条公式是 PPO 的核心。它通过“裁剪”更新幅度，避免策略一步改得太猛，从而保持训练稳定。",
        "why": "你的工作虽然重点在可塑性恢复和蒸馏，但底层优化器仍然是 PPO，因此论文必须清楚说明策略网络是怎么训练的。",
        "symbols": [
            ("L_actor", "策略网络的损失。训练时希望它尽量小。"),
            ("E[·]", "对一个 batch 的平均。"),
            ("r_t", "新旧策略概率比，r_t = exp(log π_new - log π_old)。"),
            ("A_t", "优势函数，表示当前动作比平均水平好多少。"),
            ("clip(·)", "裁剪函数，把 r_t 限制在 [1-ε, 1+ε] 内。"),
            ("ε", "PPO 裁剪范围，对应代码里的 clip_range。"),
            ("min(·,·)", "取更保守的那一项，防止策略更新过头。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "499-507",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 499, 507),
            }
        ],
    },
    {
        "title": "公式4 价值网络损失",
        "formula": "L_critic = MSE(V(s_t), R_t)",
        "role": "价值网络负责估计状态值。它的目标是让当前估计 V(s_t) 尽量接近目标回报 R_t。",
        "why": "Actor 负责“怎么做”，Critic 负责“这样做值不值”。如果 Critic 学不好，TD 误差、优势函数和后续可塑性信号都会不可靠。",
        "symbols": [
            ("L_critic", "价值网络损失。"),
            ("MSE", "均方误差，预测值和目标值之间的平方差平均。"),
            ("V(s_t)", "Critic 对当前状态的价值估计。"),
            ("R_t", "用于监督 Critic 的目标回报。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "503-507",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 503, 507),
            }
        ],
    },
    {
        "title": "公式5 双向KL蒸馏损失",
        "formula": "L_DKL = α_t D_KL(π_tem || π_θ) + (1 - α_t) D_KL(π_θ || π_tem)",
        "role": "这条公式用于“重置后把新网络拉回到旧知识附近”。前向 KL 更强调覆盖旧策略，反向 KL 更强调让当前策略不要偏离得太散。",
        "why": "Stay Hungry / P3O 的核心思想就是 Reset 后进行 Distill。你的工作则进一步把这里的 α_t 做成动态可调，因此这条公式是所有创新的中心连接点。",
        "symbols": [
            ("L_DKL", "双向 KL 蒸馏损失。"),
            ("α_t", "当前时刻的蒸馏权重，用来平衡前向 KL 和反向 KL。"),
            ("D_KL(π_tem || π_θ)", "前向 KL，要求新策略覆盖教师策略。"),
            ("D_KL(π_θ || π_tem)", "反向 KL，约束新策略不要跑得太远。"),
            ("π_tem", "临时教师策略，通常是 reset 前复制出来的旧策略。"),
            ("π_θ", "当前正在更新的策略。"),
            ("1 - α_t", "反向 KL 的权重，保证两项总权重为 1。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "372-375",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 372, 375),
            }
        ],
    },
    {
        "title": "公式6 进度驱动的动态蒸馏权重",
        "formula": "α_prog(p_t) = α_start * (1 - tanh(λ_α (p_t - 0.5))) / 2 + α_end * (1 + tanh(λ_α (p_t - 0.5))) / 2",
        "role": "这条公式让 α 随训练进度平滑变化。训练前期更靠近 α_start，后期更靠近 α_end，中间通过 tanh 做 S 型平滑过渡。",
        "why": "固定 α 的问题是“前期和后期用同一套蒸馏强度”。而训练早期更需要保守地保持旧知识，后期则更需要给新策略留出适应空间，所以要让 α 随进度变化。",
        "symbols": [
            ("α_prog(p_t)", "仅由训练进度决定的基础 α。"),
            ("p_t", "标准化训练进度，通常是 total_steps / max_steps，范围在 0 到 1。"),
            ("α_start", "训练前期采用的蒸馏权重。"),
            ("α_end", "训练后期采用的蒸馏权重。"),
            ("λ_α", "控制曲线陡峭程度的参数；越大，转折越快。"),
            ("tanh(·)", "双曲正切函数，用来实现平滑 S 型过渡。"),
            ("0.5", "把曲线转折点放在训练中期。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "378-381",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 378, 381),
            }
        ],
    },
    {
        "title": "公式7 TD反馈修正后的蒸馏权重",
        "formula": "α_base,t = clip(α_prog(p_t) - k_α (\\bar{δ}_t / (\\tilde{δ}_t + ε) - 1), α_min, α_max)",
        "role": "这条公式让 α 不仅看进度，还看“当前学习难度”。如果当前 TD 误差明显高于参考水平，说明模型还不够会，就把 α 调低，给新策略更多学习自由。",
        "why": "只看进度仍然有局限：同样训练到 50%，不同任务的学习状态可能完全不一样。引入 TD 反馈后，α 会对真实训练状态更敏感。",
        "symbols": [
            ("α_base,t", "结合进度和 TD 反馈后的基础 α。"),
            ("clip(x, α_min, α_max)", "把 α 限制在合法范围内，避免过大或过小。"),
            ("α_prog(p_t)", "进度驱动的基础 α。"),
            ("k_α", "TD 反馈对 α 的调节强度。"),
            ("\\bar{δ}_t", "当前 TD 误差的指数滑动平均。"),
            ("\\tilde{δ}_t", "参考 TD 误差的指数滑动平均。"),
            ("ε", "稳定项，防止分母过小。"),
            ("- 1", "把“比值等于 1”解释成“与参考水平相同”，便于判断偏高还是偏低。"),
            ("α_min, α_max", "α 的上下界。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "384-388",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 384, 388),
            }
        ],
    },
    {
        "title": "公式8 KL闭环更新项",
        "formula": "α_kl,t = clip(α_kl,t-1 + η_α (KL_t - κ) / max(κ, ε), α_min, α_max)",
        "role": "这条公式根据“当前策略偏离教师策略有多大”来修正 α。如果 KL 偏差过大，就把 α 提高，让蒸馏更强；如果 KL 偏差不大，α 就不会被过度抬高。",
        "why": "只看 TD 误差会偏向价值学习视角，但蒸馏本身还关心策略分布差异，所以引入 KL 闭环可以让 α 直接受策略偏移控制。",
        "symbols": [
            ("α_kl,t", "由 KL 反馈得到的 α 分量。"),
            ("α_kl,t-1", "上一时刻的 KL 闭环 α。"),
            ("η_α", "KL 闭环的更新步长。"),
            ("KL_t", "当前观察到的新旧策略 KL 偏差。"),
            ("κ", "目标 KL 水平。"),
            ("max(κ, ε)", "用于归一化，避免 κ 太小导致分母不稳定。"),
            ("α_min, α_max", "α 的取值边界。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "391-396",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 391, 396),
            }
        ],
    },
    {
        "title": "公式9 最终蒸馏权重融合",
        "formula": "α_t = β_α α_base,t + (1 - β_α) α_kl,t",
        "role": "这条公式把“进度+TD 路径”得到的 α，和“KL 闭环路径”得到的 α 融合，形成最终真正用于蒸馏的 α_t。",
        "why": "一个信号通常不够稳。进度项更平滑，KL 项更灵敏，把两者融合能兼顾趋势性与即时性。",
        "symbols": [
            ("α_t", "最终实际使用的蒸馏权重。"),
            ("β_α", "融合系数；越大，越信任 α_base,t。"),
            ("α_base,t", "进度+TD 修正得到的 α。"),
            ("α_kl,t", "KL 闭环修正得到的 α。"),
            ("1 - β_α", "KL 路径所占权重。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "399-402",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 399, 402),
            }
        ],
    },
    {
        "title": "公式10 多信号可塑性触发强度",
        "formula": "g_t = w_δ (\\bar{δ}_t/(\\tilde{δ}_t + ε) - 1) + w_KL (\\bar{KL}_t/(\\tilde{KL}_t + ε) - 1) + w_H (\\tilde{H}_t/(\\bar{H}_t + ε) - 1)",
        "role": "这条公式综合 TD 误差、策略偏移和策略熵三个信号，判断当前是否真的需要更强的可塑性恢复。",
        "why": "仅用一个信号容易误判。比如 TD 高可能只是暂时波动，但如果同时 KL 高、熵低，就更像是“模型变僵了，需要重启塑性”。",
        "symbols": [
            ("g_t", "综合塑性触发强度。"),
            ("w_δ", "TD 信号的权重。"),
            ("\\bar{δ}_t", "当前 TD 误差统计。"),
            ("\\tilde{δ}_t", "参考 TD 误差统计。"),
            ("w_KL", "KL 信号的权重。"),
            ("\\bar{KL}_t", "当前 KL 统计。"),
            ("\\tilde{KL}_t", "参考 KL 统计。"),
            ("w_H", "熵信号的权重。"),
            ("\\bar{H}_t", "当前策略熵统计。"),
            ("\\tilde{H}_t", "参考策略熵统计。"),
            ("ε", "稳定项。"),
            ("-1", "把“等于参考水平”对齐到 0，便于直接解释成偏高还是偏低。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "541-550",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 541, 550),
            }
        ],
    },
    {
        "title": "公式11 自适应重置率",
        "formula": "ρ_t = clip(ρ_0 (1 + k_r g_t), ρ_min, ρ_max)",
        "role": "这条公式让重置率不是固定常数，而是根据当前训练状态动态变化。可塑性问题越明显，ρ_t 越大；训练比较平稳时，ρ_t 就收缩。",
        "why": "固定重置率的问题是可能“该重置时不够，不该重置时太多”。引入 g_t 后，重置就能更贴合训练状态。",
        "symbols": [
            ("ρ_t", "当前时刻实际使用的重置率。"),
            ("ρ_0", "基础重置率。"),
            ("k_r", "重置率对触发强度 g_t 的敏感系数。"),
            ("g_t", "综合塑性触发强度。"),
            ("ρ_min, ρ_max", "重置率上下界。"),
            ("clip", "限制重置率，防止失控。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "527-538",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 527, 538),
            }
        ],
    },
    {
        "title": "公式12 神经元重要性评分",
        "formula": "I_j = w_a mean(|h_j|) + w_w mean(|W_next[:, j]|)",
        "role": "这条公式用于判断某个隐藏神经元“重要不重要”。激活大、对下一层影响强的神经元，更不应该被优先重置。",
        "why": "如果 reset 完全随机，可能把关键知识也一起打掉。做重要性评分后，可以优先重置低重要性神经元，提高 reset 的精度。",
        "symbols": [
            ("I_j", "第 j 个隐藏神经元的重要性分数。"),
            ("w_a", "激活项权重。"),
            ("mean(|h_j|)", "该神经元在一批状态上的平均绝对激活。"),
            ("h_j", "第 j 个神经元的激活值。"),
            ("w_w", "出边权重项权重。"),
            ("W_next[:, j]", "该神经元连向下一层的权重列。"),
            ("mean(|W_next[:, j]|)", "该神经元对下一层平均影响强度。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "205-232",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 205, 232),
            }
        ],
    },
    {
        "title": "公式13 选择性重置",
        "formula": "J_reset = arg TopLow( I_j, K ),  仅重置低重要性神经元",
        "role": "这不是传统教科书里的标准公式，但它对应了你论文里一个很重要的实现思想：先用 I_j 排序，再优先重置重要性最低的那部分神经元。",
        "why": "相比随机重置，选择性重置更像“有保留地恢复塑性”：把不太关键的单元刷新掉，同时保留重要单元中的已有知识。",
        "symbols": [
            ("J_reset", "被选中进行重置的神经元索引集合。"),
            ("TopLow", "从低到高排序后取前 K 个。"),
            ("I_j", "神经元重要性分数。"),
            ("K", "本轮需要重置的神经元数量。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "261-298",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 261, 298),
            }
        ],
    },
    {
        "title": "公式14 困难状态优先级",
        "formula": "q(s) = [ w_td · δ_norm(s) + w_kl · KL_norm(s) ]^γ_p",
        "role": "这条公式给每个状态打一个“优先蒸馏分数”。TD 误差越大、KL 偏差越大，说明这个状态越难、越重要，就越该优先参与蒸馏。",
        "why": "随机采样的问题是把关键状态和普通状态一视同仁。优先级采样能把蒸馏资源集中在更容易遗忘、也更影响性能的状态上。",
        "symbols": [
            ("q(s)", "状态 s 的优先级分数。"),
            ("w_td", "TD 项权重。"),
            ("δ_norm(s)", "状态 s 的归一化 TD 误差。"),
            ("w_kl", "KL 项权重。"),
            ("KL_norm(s)", "状态 s 的归一化策略偏差。"),
            ("γ_p", "优先级放大指数，对应代码中的 priority_gamma。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "423-440",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 423, 440),
            }
        ],
    },
    {
        "title": "公式15 困难状态集合",
        "formula": "B_hard = TopK_{s ∈ B}( q(s), K )",
        "role": "从候选状态池 B 中，选出优先级最高的 K 个状态，组成困难状态集合。",
        "why": "论文里写成 TopK 的形式是为了说明“优先挑最重要的状态”；而代码里为了更平滑，采用了按 priority 概率采样的实现。",
        "symbols": [
            ("B_hard", "困难状态集合。"),
            ("TopK", "取分数最高的前 K 个。"),
            ("s ∈ B", "状态 s 属于候选缓冲区 B。"),
            ("q(s)", "状态优先级分数。"),
            ("K", "选取的状态数量。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "327-355",
                "intro": "priority 采样版本的实现如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 327, 355),
            },
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "361-369",
                "intro": "在不使用 priority 时，代码也保留了按高 TD 状态取 top pool 的近似实现：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 361, 369),
            },
        ],
    },
    {
        "title": "公式16 最终蒸馏状态批次",
        "formula": "B_distill = B_hard ∪ B_rand",
        "role": "最终进入蒸馏的状态，不只包含困难状态，也混入一部分随机状态，避免只盯着少量极端样本。",
        "why": "如果只蒸馏 hardest states，容易覆盖不全面；加入随机样本后，既保留重点，又保持整体代表性。",
        "symbols": [
            ("B_distill", "最终蒸馏 batch。"),
            ("B_hard", "优先挑出的困难状态。"),
            ("∪", "并集，表示两部分状态合并。"),
            ("B_rand", "随机补充状态。"),
        ],
        "code_refs": [
            {
                "path": "run_p3o_extra_innovations.py",
                "lines": "340-354",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_p3o_extra_innovations.py", 340, 354),
            }
        ],
    },
    {
        "title": "公式17 锚点增强总蒸馏损失",
        "formula": "L_total = L_current + λ_anchor L_anchor",
        "role": "这条公式把当前任务蒸馏损失和跨阶段锚点恢复损失加在一起。它的目标是：既学当前，又别忘旧任务。",
        "why": "如果只对当前阶段做蒸馏，模型可能仍然遗忘更早阶段的知识；加入 anchor loss 后，相当于给更早的关键状态加了一层记忆保护。",
        "symbols": [
            ("L_total", "最终优化的总蒸馏损失。"),
            ("L_current", "当前阶段的蒸馏损失。"),
            ("λ_anchor", "锚点恢复损失权重，对应代码中的 anchor_loss_weight。"),
            ("L_anchor", "锚点恢复损失。"),
        ],
        "code_refs": [
            {
                "path": "run_sequence_plasticity.py",
                "lines": "593-617",
                "intro": "对应代码如下：",
                "snippet": get_snippet("run_sequence_plasticity.py", 593, 617),
            }
        ],
    },
    {
        "title": "公式18 锚点恢复损失",
        "formula": "L_anchor = E_{s ∼ M}[ q(s) · ( α_t D_KL(π_ref || π_θ) + (1-α_t) D_KL(π_θ || π_ref) ) ]",
        "role": "这条公式表示：从锚点记忆库 M 中取出历史关键状态，用加权双向蒸馏把当前策略重新拉向旧任务参考策略。",
        "why": "它比普通 Distill 更强的地方在于：不是只记住“上一刻的老师”，而是把跨阶段保留下来的关键状态也纳入训练。",
        "symbols": [
            ("L_anchor", "锚点恢复损失。"),
            ("E_{s ∼ M}[·]", "对锚点记忆库 M 中采样状态求平均。"),
            ("M", "锚点记忆库。"),
            ("q(s)", "状态权重，优先强调更关键的锚点状态。"),
            ("α_t", "当前蒸馏权重。"),
            ("π_ref", "锚点保存的参考教师策略。"),
            ("π_θ", "当前策略。"),
            ("1-α_t", "反向 KL 的权重。"),
        ],
        "code_refs": [
            {
                "path": "run_sequence_plasticity.py",
                "lines": "552-590",
                "intro": "锚点候选构建如下：",
                "snippet": get_snippet("run_sequence_plasticity.py", 552, 590),
            },
            {
                "path": "run_sequence_plasticity.py",
                "lines": "607-617",
                "intro": "锚点损失实际接入训练如下：",
                "snippet": get_snippet("run_sequence_plasticity.py", 607, 617),
            },
            {
                "path": "run_sequence_plasticity.py",
                "lines": "542-549",
                "intro": "锚点部分最终仍然调用加权双向 KL：",
                "snippet": get_snippet("run_sequence_plasticity.py", 542, 549),
            },
        ],
    },
]


def build_doc():
    doc = Document()

    add_para(doc, "论文全部核心公式详解与代码对应", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(
        doc,
        "说明：本文件按“公式 -> 作用 -> 为什么这样写 -> 每个符号解释 -> 对应代码”整理，便于你写论文、答辩和向老师逐条说明。代码片段均来自当前项目中的真实实现文件。",
        size=11,
    )

    for item in FORMULAS:
        add_para(doc, item["title"], size=14, bold=True)
        add_para(doc, f"公式：{item['formula']}", size=11, bold=True)
        add_para(doc, f"这条公式做什么：{item['role']}", size=11)
        add_para(doc, f"为什么这样写：{item['why']}", size=11)
        add_para(doc, "符号详细解释：", size=11, bold=True)
        for symbol, explanation in item["symbols"]:
            add_para(doc, f"{symbol}：{explanation}", size=11)

        for idx, ref in enumerate(item["code_refs"], start=1):
            add_para(
                doc,
                f"代码对应 {idx}：`{ref['path']}:{ref['lines']}`",
                size=11,
                bold=True,
            )
            add_para(doc, ref["intro"], size=11)
            add_code_block(doc, ref["snippet"])

    add_para(doc, "补充说明", size=14, bold=True)
    add_para(doc, "1. 这份文档覆盖了论文方法部分和实验部分真正用到的全部核心公式，包括 PPO 基础、动态 α、自适应重置、困难状态蒸馏和锚点恢复。", size=11)
    add_para(doc, "2. 某些论文公式是理论表达，代码里不一定逐字逐句写成完全相同的数学形式，但功能是一一对应的。", size=11)
    add_para(doc, "3. 如果你后面还需要，我可以继续给你做一版“只保留公式+符号解释，不带代码”的精简文档，适合直接贴入论文附录。", size=11)

    doc.save(str(OUTPUT))
    print(f"saved: {OUTPUT}")


if __name__ == "__main__":
    build_doc()
